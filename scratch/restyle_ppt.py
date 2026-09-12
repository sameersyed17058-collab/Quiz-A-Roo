from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

SOURCE = 'Quiz-A-Roo (1).pptx'
OUTPUT = 'Quiz-A-Roo (1)-project-aligned.pptx'

LIGHT = RGBColor(0xFC, 0xFB, 0xF9)
DIM = RGBColor(0xD6, 0xD3, 0xD1)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
GOLD = RGBColor(0xFB, 0xBF, 0x24)
INK = RGBColor(0x0C, 0x0A, 0x09)

replacements = {
    'AI-Powered Adaptive Quiz Platform | Group Members: [Name 1, Name 2, Name 3, etc.]':
        'Interactive AI Outback Quest | Adaptive trivia, gamified progression, and reliable fallback questions',
    'Google Stitch': 'React.js + Stitch-inspired UI',
    'A resilient, end-to-end pipeline ensures every player receives a fair, personalised quiz experience — with or without live AI access.':
        'A resilient pipeline gives every Explorer a fair, personalized quiz experience, with or without live AI access.',
    'Privacy preserved — Secret Explorer IDs keep identities safe':
        'Privacy preserved with password-free Secret Explorer IDs',
    'Every UI decision prioritises simplicity, speed, and player delight.':
        'Every UI decision prioritizes simplicity, speed, and player delight.',
}

def replace_text(text):
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text

def set_run_style(run, heading=False):
    run.font.name = 'Outfit' if heading else 'Plus Jakarta Sans'
    run.font.bold = heading or run.font.bold
    run.font.color.rgb = LIGHT if heading else DIM
    if heading:
        run.font.color.rgb = GOLD

presentation = Presentation(SOURCE)

for slide in presentation.slides:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = INK

    for shape in slide.shapes:
        if not hasattr(shape, 'text_frame'):
            continue
        original = shape.text
        updated = replace_text(original)
        if updated != original:
            shape.text = updated

        is_heading = any(run.font.name == 'Comfortaa Bold' for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                set_run_style(run, heading=is_heading)
                if 'Quiz-A-Roo' in run.text or run.text.strip() in {'COMPETE. IMPROVE. RISE.', 'Thank You!'}:
                    run.font.color.rgb = AMBER

presentation.save(OUTPUT)
print(OUTPUT)